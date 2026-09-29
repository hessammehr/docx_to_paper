# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1"]
# ///
"""Test filters/docx-colors.lua on a generated .docx.

Run from the repository root:  uv run tests/docx-colors-test.py
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parent.parent
FILTER = ROOT / "filters" / "docx-colors.lua"


def run(par, text, color=None, fill=None, bold=False, italic=False):
    r = par.add_run(text)
    r.bold, r.italic = bold or None, italic or None
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        r._r.get_or_add_rPr().append(shd)
    return r


def build(path):
    doc = Document()

    # 1. Timeline: coloured bold label inside a shaded stretch; the label's
    #    bold run continues into "W0." (pandoc merges them into one Strong).
    p = doc.add_paragraph()
    run(p, "Samples:", "BF4E14", "FAE2D5", bold=True)
    run(p, " W0.", None, "FAE2D5", bold=True)
    run(p, " Samples ordered.", None, "FAE2D5")
    run(p, " ")
    run(p, "Data:", "0B769F", "CAEDFB", bold=True)
    run(p, " ", None, "CAEDFB")
    run(p, "W1.", None, "CAEDFB", bold=True)
    run(p, " Storage set up.", None, "CAEDFB")

    # 2. Colour starting mid-word and crossing out of italics.
    p = doc.add_paragraph()
    run(p, "An un")
    run(p, "believ", "C00000")
    run(p, "able ")
    run(p, "alpha ", italic=True)
    run(p, "beta", "2E75B6", italic=True)
    run(p, " gamma", "2E75B6")
    run(p, " delta.")

    # 3. Shading crossing out of italics: the space after the italics must be
    #    shaded too (no gap).
    p = doc.add_paragraph()
    run(p, "one ", italic=True)
    run(p, "two", None, "FFF5CC", italic=True)
    run(p, " three", None, "FFF5CC")
    run(p, " four.")

    # 4. Noise that should be ignored.
    p = doc.add_paragraph()
    run(p, "Stray")
    run(p, ", ", "196B24")
    run(p, "near-black ", "222222")
    run(p, "and white shading.", None, "FFFFFF")

    # 5. Colour inside a table cell (pandoc: Plain block).
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    run(cell.paragraphs[0], "Cell ")
    run(cell.paragraphs[0], "red", "FF0000")

    doc.save(path)


EXPECTED = [
    '[**[Samples:]{color="#BF4E14"} W0.** Samples ordered.]{shading="#FAE2D5"}'
    ' [[**Data:**]{color="#0B769F"} **W1.** Storage set up.]{shading="#CAEDFB"}',
    'An un[believ]{color="#C00000"}able *alpha [beta]{color="#2E75B6"}*'
    '[ gamma]{color="#2E75B6"} delta.',
    '*one [two]{shading="#FFF5CC"}*[ three]{shading="#FFF5CC"} four.',
    "Stray, near-black and white shading.",
    '[red]{color="#FF0000"}',
]


def main():
    with tempfile.TemporaryDirectory() as tmp:
        docx = Path(tmp) / "colors.docx"
        build(docx)
        res = subprocess.run(
            ["pandoc", "-f", "docx", "--lua-filter", str(FILTER), str(docx),
             "-t", "markdown", "--wrap=none"],
            capture_output=True, text=True, check=True,
        )
    out = res.stdout
    failed = [e for e in EXPECTED if e not in out]
    if res.stderr.strip():
        print("stderr:\n" + res.stderr)
    if failed:
        print("FAIL: missing from output:\n  " + "\n  ".join(failed))
        print("\nOutput:\n" + out)
        sys.exit(1)
    print(f"OK: {len(EXPECTED)} checks passed")


if __name__ == "__main__":
    main()
